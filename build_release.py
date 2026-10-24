"""
Assemble the folder to upload to Mendeley Data (HMT-S), from the correct sources:

  Tier 1 raw EEG   <- Dataset/2. HMT-S/1. Raw EEG and Stimulus Logs/<class>/Subject N/*.csv
  Tier 1 logs      <- Dataset/2. HMT-S/5. Stimulus Logs (anonymized)/<class>/Subject N/LS-Quiz-subjectN.xlsx
  Tier 2           <- Dataset/2. HMT-S/2.|3. Preprocessed-*/   (original files)
                      except the files listed as "rebuilt" in rebuilt_tier2/rebuild_report.csv,
                      which are taken from rebuilt_tier2/
  README.md / .pdf <- Paper/SSRN Dataset/README_HMT-S.md (and README_HMT-S.pdf if present)

Nothing in the dataset folder is changed. The release folder is created next to the
dataset folder and must not exist yet.

Checks before anything is copied (the script stops if one fails):
  - 28 recordings, each with one raw CSV, one anonymized log and one Tier 2 file
  - every log sheet is named "LS-Quiz-subjectN" and has no author metadata
  - cell contents of every anonymized log equal the original log
  - every rebuilt Tier 2 file exists

Run:  python build_release.py
Then: zip the release folder and upload it; keep release_checksums.csv for your records.
"""
import hashlib
import re
import shutil
import sys

import openpyxl
import pandas as pd

import config as C

ANON_DIR = C.DATASET_DIR / "5. Stimulus Logs (anonymized)"
REBUILT_DIR = C.HERE / "rebuilt_tier2"
README_SRC = C.HERE.parent / "README_HMT-S.md"
RELEASE_DIR = C.DATASET_DIR.parent / "HMT-S Mendeley upload"
CHECKSUMS = C.HERE / "release_checksums.csv"

RAW_SUB = "1. Raw EEG and Stimulus Logs"
TIER2_SUB = {"active": "3. Preprocessed-active-reflective", "reflective": "3. Preprocessed-active-reflective",
             "verbal": "2. Preprocessed-verbal-visual", "visual": "2. Preprocessed-verbal-visual"}
PERSONAL_HINT = re.compile(r"hmt_history", re.I)


def one(folder, pattern):
    hits = [p for p in folder.iterdir() if p.is_file() and re.fullmatch(pattern, p.name, flags=re.I)] if folder.exists() else []
    return hits


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def main():
    problems, plan = [], []          # plan: (source, destination relative to RELEASE_DIR)
    if RELEASE_DIR.exists():
        sys.exit(f"{RELEASE_DIR} already exists. Delete or rename it first.")
    rebuilt = set()
    rep = REBUILT_DIR / "rebuild_report.csv"
    if rep.exists():
        r = pd.read_csv(rep)
        rebuilt = set(r.loc[r["action"] == "rebuilt", "file"])
    else:
        problems.append("rebuilt_tier2/rebuild_report.csv not found (run rebuild_tier2.py first)")

    n_rec = 0
    for cls, folder in C.CLASS_FOLDERS.items():
        for sd in sorted((C.DATASET_DIR / RAW_SUB / folder).glob("Subject *")):
            num = int(sd.name.split()[-1])
            n_rec += 1
            tag = f"{cls} subject {num}"
            # raw EEG
            raws = one(sd, r".*\.csv")
            if len(raws) != 1:
                problems.append(f"{tag}: expected one raw .csv, found {[p.name for p in raws]}")
            else:
                plan.append((raws[0], f"{RAW_SUB}/{folder}/Subject {num}/Subject{num}.csv"))
            # anonymized log
            logs = one(ANON_DIR / folder / f"Subject {num}", r"ls-quiz.*\.xlsx")
            if len(logs) != 1:
                problems.append(f"{tag}: expected one anonymized log, found {[p.name for p in logs]}")
            else:
                lg = logs[0]
                wb = openpyxl.load_workbook(lg, read_only=True)
                if wb.sheetnames != [f"LS-Quiz-subject{num}"] or any(PERSONAL_HINT.search(s) for s in wb.sheetnames):
                    problems.append(f"{tag}: log sheet name {wb.sheetnames} is not anonymized")
                if wb.properties.lastModifiedBy or (wb.properties.creator not in (None, "", "openpyxl")):
                    problems.append(f"{tag}: log metadata not empty (creator={wb.properties.creator}, "
                                    f"lastModifiedBy={wb.properties.lastModifiedBy})")
                wb.close()
                orig = one(sd, r"ls-quiz.*\.xlsx")
                if orig and not pd.read_excel(orig[0]).equals(pd.read_excel(lg)):
                    problems.append(f"{tag}: anonymized log content differs from the original log")
                plan.append((lg, f"{RAW_SUB}/{folder}/Subject {num}/LS-Quiz-subject{num}.xlsx"))
            # Tier 2
            name = f"{cls}-{num:02d}-subject{num}-preprocessed.csv"
            src = (REBUILT_DIR if name in rebuilt else C.DATASET_DIR) / TIER2_SUB[cls] / name
            if not src.exists():
                problems.append(f"{tag}: Tier 2 file not found: {src}")
            else:
                plan.append((src, f"{TIER2_SUB[cls]}/{name}"))

    if n_rec != 28:
        problems.append(f"expected 28 recordings, found {n_rec}")
    if not README_SRC.exists():
        problems.append(f"README not found: {README_SRC}")
    else:
        plan.append((README_SRC, "README.md"))
        pdf = README_SRC.with_suffix(".pdf")
        if pdf.exists():
            plan.append((pdf, "README.pdf"))
        else:
            print("Note: README_HMT-S.pdf not found; only README.md will be included.")

    if problems:
        print("STOPPED - nothing was copied:\n  " + "\n  ".join(problems))
        sys.exit(1)

    rows = []
    for src, rel in plan:
        dst = RELEASE_DIR / rel
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
        rows.append(dict(file=rel, source=str(src), bytes=dst.stat().st_size, sha256=sha256(dst)))
        print(f"  {rel}")
    pd.DataFrame(rows).to_csv(CHECKSUMS, index=False)
    total = sum(r["bytes"] for r in rows) / 1e6
    print(f"\n{len(rows)} files, {total:.1f} MB -> {RELEASE_DIR}")
    print(f"Tier 2 from rebuilt_tier2/: {len(rebuilt)} files")
    print(f"Checksums: {CHECKSUMS}")


if __name__ == "__main__":
    main()
